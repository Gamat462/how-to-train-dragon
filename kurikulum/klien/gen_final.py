# Kode yang akan dimuat di halaman: membuat berkas klien tiruan.
import numpy as np, pandas as pd
from openpyxl import Workbook

rng = np.random.default_rng(7)
cabang = ["Kemang", "Dago", "Seturan"]
harga  = {"Kopi Susu":22000, "Americano":18000, "Roti Bakar":15000,
          "Croissant":25000, "Teh Manis":10000}
produk = list(harga)

def baris(n, b1, b2):
    out = []
    for _ in range(n):
        b = int(rng.integers(b1, b2+1)); h = int(rng.integers(1, 29))
        p = produk[int(rng.integers(0, 5))]; c = cabang[int(rng.integers(0, 3))]
        q = int(rng.integers(1, 6))
        out.append({"Tanggal": f"{h:02d}/{b:02d}/2024", "Cabang": c, "Produk": p,
                    "Jumlah": q, "Harga Satuan": harga[p], "Total": q*harga[p]})
    return out

wb = Workbook(); ws = wb.active; ws.title = "Penjualan Q1"
ws["A1"] = "LAPORAN PENJUALAN KEDAI KOPI SENJA"
ws["A2"] = "Periode: Januari - Maret 2024"
kol = ["Tanggal","Cabang","Produk","Jumlah","Harga Satuan","Total"]
for i, k in enumerate(kol, 1): ws.cell(row=4, column=i, value=k)

d = baris(180, 1, 3)
d[10]["Cabang"] = " kemang "; d[11]["Cabang"] = "KEMANG"; d[12]["Cabang"] = "Kemang "
d[40]["Jumlah"] = None; d[41]["Total"] = None
d[70]["Harga Satuan"] = "22.000"; d[71]["Harga Satuan"] = "18.000"; d[72]["Total"] = "45.000"
d[100]["Tanggal"] = "2024-02-15"; d[101]["Tanggal"] = "17 Feb 2024"
d.append(dict(d[5])); d.append(dict(d[5]))
for r, row in enumerate(d, 5):
    for c, k in enumerate(kol, 1): ws.cell(row=r, column=c, value=row[k])
ws.cell(row=len(d)+6, column=1, value="TOTAL")

ws2 = wb.create_sheet("Penjualan Q2")
kol2 = ["TANGGAL","CABANG","PRODUK","QTY","HARGA","TOTAL"]
for i, k in enumerate(kol2, 1): ws2.cell(row=1, column=i, value=k)
peta = dict(zip(kol2, kol))
for r, row in enumerate(baris(150, 4, 6), 2):
    for c, k in enumerate(kol2, 1): ws2.cell(row=r, column=c, value=row[peta[k]])

wb.create_sheet("Catatan")["A1"] = "Jangan diubah. Hubungi Rina kalau ada pertanyaan."
wb.save("Laporan Penjualan 2024.xlsx")
print("Berkas 'Laporan Penjualan 2024.xlsx' dibuat.")

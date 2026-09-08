"""Kasus 1 — Toko Bangunan Sumber Rejeki.
Meniru ekspor Excel dari aplikasi kasir murah: judul di sel gabungan,
header bukan di baris 1, angka jadi teks, tanggal tiga format,
baris subtotal nyempil, dan footer cetak di bawah."""
import random
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

rng = random.Random(2024)

PRODUK = [
    ("SMN-40KG",  "Semen Gresik 40kg",        62000,  48000),
    ("SMN 40KG",  "Semen Gresik 40 kg",       62000,  48000),   # kode sama, tulisan beda
    ("BTA-MRH",   "Bata Merah (per 100)",     95000,  72000),
    ("PSR-M3",    "Pasir Beton (m3)",        320000, 245000),
    ("CAT-5KG",   "Cat Tembok Avian 5kg",    185000, 142000),
    ("cat-5kg",   "Cat tembok avian 5 kg",   185000, 142000),   # kode huruf kecil
    ("PPA-4IN",   "Pipa PVC 4 inci",          78000,  59000),
    ("KRM-60",    "Keramik 60x60 (dus)",     145000, 112000),
    ("PKU-KG",    "Paku 5cm (kg)",            24000,  17000),
]

PELANGGAN = [
    "Budi Santoso", "budi santoso", "BUDI SANTOSO", "Budi  Santoso",  # satu orang, empat tulisan
    "CV Karya Abadi", "Cv Karya Abadi", "CV. Karya Abadi",
    "Toko Berkah", "toko berkah",
    "Hj. Aminah", "Pak Slamet", "Ibu Ratna", "PT Mitra Bangun",
    "Umum", "umum", "-",
]

def fmt_tanggal(d, gaya):
    if gaya == 0: return d.strftime("%d/%m/%Y")      # 01/03/2024
    if gaya == 1: return d.strftime("%Y-%m-%d")      # 2024-03-01
    return d.strftime("%d-%b-%y")                     # 01-Mar-24

def rupiah_teks(n):
    return f"{n:,}".replace(",", ".")                 # 1250000 -> "1.250.000"

wb = Workbook()
ws = wb.active
ws.title = "Sheet1"

# --- judul di sel gabungan, header baru di baris 6 ---
ws.merge_cells("B1:H1"); ws["B1"] = "TOKO BANGUNAN SUMBER REJEKI"
ws["B1"].font = Font(bold=True, size=14); ws["B1"].alignment = Alignment(horizontal="center")
ws.merge_cells("B2:H2"); ws["B2"] = "Jl. Raya Bogor KM 24 No. 88, Depok"
ws["B2"].alignment = Alignment(horizontal="center")
ws.merge_cells("B3:H3"); ws["B3"] = "LAPORAN PENJUALAN PERIODE JAN - JUN 2024"
ws["B3"].font = Font(bold=True); ws["B3"].alignment = Alignment(horizontal="center")
# baris 4 & 5 kosong

HEADER = ["No", "Tanggal", "No. Nota", "Kode Barang", "Nama Barang",
          "Qty", "Harga Satuan", "Jumlah", "Pelanggan"]
for i, h in enumerate(HEADER):
    c = ws.cell(row=6, column=2+i, value=h); c.font = Font(bold=True)

baris = 7
no = 0
total_baris_data = 0
for bulan in range(1, 7):
    subtotal = 0
    hari_awal = date(2024, bulan, 1)
    hari_akhir = date(2024, bulan+1, 1) - timedelta(days=1) if bulan < 12 else date(2024,12,31)
    n_trx = rng.randint(22, 30)
    for _ in range(n_trx):
        no += 1
        d = hari_awal + timedelta(days=rng.randrange((hari_akhir-hari_awal).days + 1))
        kode, nama, jual, _beli = PRODUK[rng.randrange(len(PRODUK))]
        qty = rng.choice([1,1,2,2,3,4,5,10,20])
        # ~6% retur: qty negatif, tanpa penanda apa pun
        if rng.random() < 0.06:
            qty = -rng.choice([1,2])
        jumlah = qty * jual
        subtotal += jumlah
        ws.cell(row=baris, column=2, value=no)
        ws.cell(row=baris, column=3, value=fmt_tanggal(d, rng.randrange(3)))
        ws.cell(row=baris, column=4, value=f"NT{bulan:02d}{no:04d}")
        ws.cell(row=baris, column=5, value=kode)
        ws.cell(row=baris, column=6, value=nama)
        ws.cell(row=baris, column=7, value=qty)
        ws.cell(row=baris, column=8, value=rupiah_teks(jual))     # TEKS, bukan angka
        ws.cell(row=baris, column=9, value=rupiah_teks(jumlah))   # TEKS
        ws.cell(row=baris, column=10, value=PELANGGAN[rng.randrange(len(PELANGGAN))])
        baris += 1
        total_baris_data += 1
    # baris subtotal nyempil di tengah data
    c = ws.cell(row=baris, column=6, value=f"TOTAL BULAN {bulan:02d}"); c.font = Font(bold=True)
    c2 = ws.cell(row=baris, column=9, value=rupiah_teks(subtotal)); c2.font = Font(bold=True)
    baris += 2   # + satu baris kosong pemisah

# footer cetak
ws.cell(row=baris+1, column=2, value="Dicetak oleh: Admin Kasir")
ws.cell(row=baris+2, column=2, value="Tanggal cetak: 05/07/2024")
ws.cell(row=baris+3, column=2, value="Aplikasi Kasir Pro v2.1 - Lisensi: TOKO SUMBER REJEKI")

wb.save("Laporan Penjualan Toko Sumber Rejeki.xlsx")
print("baris transaksi asli :", total_baris_data)
print("baris terakhir dipakai:", baris+3)

"""Kasus 3 — CV Karya Mandiri (kontraktor).
Dua sheet dalam satu workbook, dibuat dua orang berbeda:
  'RAB'      -> nama proyek versi marketing
  'Realisasi'-> nama proyek versi lapangan, plus biaya campur teks & angka
"""
import random
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font

rng = random.Random(31)

PROYEK = [
 # (nama di RAB,                     nama di Realisasi,          nilai kontrak, rab)
 ("Pembangunan Ruko Depok 2 Lantai", "RUKO DEPOK - 2024",        850_000_000, 690_000_000),
 ("Renovasi Kantor PT Sinar Mas",    "renovasi kantor sinarmas",  320_000_000, 268_000_000),
 ("Gudang Logistik Cibitung",        "GUDANG CIBITUNG",         1_240_000_000, 985_000_000),
 ("Rumah Tinggal Bpk. Handoko",      "Rmh Bpk Handoko",           540_000_000, 448_000_000),
 ("Pagar & Taman Perumahan Griya",   "PAGAR TAMAN GRIYA",         185_000_000, 149_000_000),
 ("Masjid Al-Ikhlas Tahap 2",        "MASJID AL IKHLAS TAHAP II", 410_000_000, 352_000_000),
 ("Kios Pasar Modern Bekasi",        "KIOS PASAR BEKASI",         275_000_000, 221_000_000),
]
KATEGORI = ["Material", "Upah Tukang", "Sewa Alat", "Transport", "Subkontraktor", "Overhead"]

wb = Workbook()

# ---------- Sheet RAB ----------
ws = wb.active; ws.title = "RAB"
ws["A1"] = "CV KARYA MANDIRI - RENCANA ANGGARAN BIAYA 2024"; ws["A1"].font = Font(bold=True, size=12)
ws.append([]); 
for i, h in enumerate(["Nama Proyek", "Nilai Kontrak", "Total RAB", "Target Margin"]):
    c = ws.cell(row=3, column=1+i, value=h); c.font = Font(bold=True)
r = 4
for nama, _, kontrak, rab in PROYEK:
    ws.cell(row=r, column=1, value=nama)
    ws.cell(row=r, column=2, value=kontrak)
    ws.cell(row=r, column=3, value=rab)
    ws.cell(row=r, column=4, value=(kontrak - rab) / kontrak)
    r += 1

# ---------- Sheet Realisasi ----------
ws2 = wb.create_sheet("Realisasi")
for i, h in enumerate(["Tanggal", "Proyek", "Kategori", "Uraian", "Biaya"]):
    c = ws2.cell(row=1, column=1+i, value=h); c.font = Font(bold=True)
r = 2
n_baris = 0
total_per_proyek = {}
for nama_rab, nama_real, kontrak, rab in PROYEK:
    # realisasi 88%-118% dari RAB -> ada yang untung, ada yang jebol
    faktor = rng.uniform(0.88, 1.18)
    target = int(rab * faktor)
    terkumpul = 0
    while terkumpul < target * 0.97:
        sisa = target - terkumpul
        biaya = min(sisa, int(rng.uniform(0.03, 0.15) * target))
        biaya = max(biaya, 500_000)
        biaya = round(biaya, -4)
        terkumpul += biaya
        d = date(2024, 1, 1) + timedelta(days=rng.randrange(210))
        kat = KATEGORI[rng.randrange(len(KATEGORI))]
        # separuh biaya ditulis sebagai TEKS "Rp 15.000.000", separuh sebagai angka
        if rng.random() < 0.5:
            nilai = f"Rp {biaya:,}".replace(",", ".")
        else:
            nilai = biaya
        ws2.cell(row=r, column=1, value=d.strftime("%d/%m/%Y"))
        ws2.cell(row=r, column=2, value=nama_real)
        ws2.cell(row=r, column=3, value=kat)
        ws2.cell(row=r, column=4, value=f"{kat} termin {rng.randint(1,4)}")
        ws2.cell(row=r, column=5, value=nilai)
        r += 1; n_baris += 1
    total_per_proyek[nama_real] = terkumpul

# satu proyek muncul di Realisasi tapi TIDAK ada di RAB (pekerjaan tambahan)
for _ in range(6):
    biaya = round(rng.randint(8_000_000, 25_000_000), -4)
    d = date(2024, 5, 1) + timedelta(days=rng.randrange(60))
    ws2.cell(row=r, column=1, value=d.strftime("%d/%m/%Y"))
    ws2.cell(row=r, column=2, value="PEK. TAMBAH KURANG RUKO DEPOK")
    ws2.cell(row=r, column=3, value="Material")
    ws2.cell(row=r, column=4, value="Tambahan permintaan owner")
    ws2.cell(row=r, column=5, value=f"Rp {biaya:,}".replace(",", "."))
    r += 1; n_baris += 1
    total_per_proyek["PEK. TAMBAH KURANG RUKO DEPOK"] = \
        total_per_proyek.get("PEK. TAMBAH KURANG RUKO DEPOK", 0) + biaya

wb.save("RAB vs Realisasi CV Karya Mandiri.xlsx")
print("proyek di RAB      :", len(PROYEK))
print("baris realisasi    :", n_baris)
print("proyek di Realisasi:", len(total_per_proyek))
for k, v in total_per_proyek.items():
    print(f"  {k:32s} {v:>14,}".replace(",", "."))

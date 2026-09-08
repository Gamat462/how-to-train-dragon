"""Kasus 4 — Yayasan Bina Cendekia.
Rekap SPP dalam bentuk MELEBAR: satu kolom per bulan.
Isinya campur: angka, teks 'LUNAS', '-', kosong, dan catatan bebas.
"""
import random
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

rng = random.Random(99)
BULAN = ["Jul","Agu","Sep","Okt","Nov","Des","Jan","Feb","Mar","Apr","Mei","Jun"]
KELAS = ["7A","7B","8A","8B","9A","9B"]
SPP   = {"7A":450000,"7B":450000,"8A":475000,"8B":475000,"9A":500000,"9B":500000}

nama_depan = ["Andi","Bella","Citra","Dimas","Eka","Farhan","Gita","Hana","Irfan","Jihan",
              "Kirana","Lukman","Mira","Naufal","Okta","Putri","Qori","Rizky","Salsa","Tegar",
              "Umar","Vina","Wahyu","Yuda","Zahra","Aditya","Bunga","Cahya","Dewa","Elsa"]
nama_blkg = ["Pratama","Wijaya","Saputra","Ramadhan","Anggraini","Nugroho","Permata","Hidayat",
             "Kusuma","Maulana"]

wb = Workbook(); ws = wb.active; ws.title = "Rekap SPP"
ws.merge_cells("A1:P1")
ws["A1"] = "YAYASAN BINA CENDEKIA — REKAP PEMBAYARAN SPP TP 2023/2024"
ws["A1"].font = Font(bold=True, size=12); ws["A1"].alignment = Alignment(horizontal="center")

HEAD = ["NIS","Nama Siswa","Kelas","SPP/Bulan"] + BULAN + ["Keterangan"]
for i,h in enumerate(HEAD):
    c = ws.cell(row=3, column=1+i, value=h); c.font = Font(bold=True)

r = 4
n_siswa = 0
n_lunas = 0
fakta = []      # kebenaran yang kita simpan untuk memverifikasi jawaban nanti
for i in range(72):
    nis = f"2023{i+1:03d}"
    nama = f"{nama_depan[i % len(nama_depan)]} {nama_blkg[rng.randrange(len(nama_blkg))]}"
    kelas = KELAS[i % len(KELAS)]
    spp = SPP[kelas]
    # 7% siswa keluar di tengah tahun
    keluar_di = rng.randrange(4, 11) if rng.random() < 0.07 else None
    ket = ""
    if keluar_di is not None:
        ket = f"Pindah sekolah bulan {BULAN[keluar_di]}"
    ws.cell(row=r, column=1, value=nis)
    ws.cell(row=r, column=2, value=nama)
    ws.cell(row=r, column=3, value=kelas)
    ws.cell(row=r, column=4, value=spp)
    bayar_total = 0
    bayar_angka = 0
    tunggak_bulan = 0
    for j, b in enumerate(BULAN):
        kol = 5 + j
        if keluar_di is not None and j >= keluar_di:
            ws.cell(row=r, column=kol, value="-")          # sudah keluar
            continue
        p = rng.random()
        if p < 0.72:                                        # bayar penuh, ditulis angka
            ws.cell(row=r, column=kol, value=spp); bayar_total += spp; bayar_angka += spp
        elif p < 0.84:                                      # bayar penuh, ditulis 'LUNAS'
            ws.cell(row=r, column=kol, value="LUNAS"); bayar_total += spp; n_lunas += 1
        elif p < 0.90:                                      # cicil sebagian
            sebagian = int(round(spp * rng.choice([0.4,0.5,0.6]), -3))
            ws.cell(row=r, column=kol, value=sebagian); bayar_total += sebagian; bayar_angka += sebagian
            tunggak_bulan += 1
        elif p < 0.96:                                      # belum bayar, dikosongkan
            tunggak_bulan += 1
        else:                                               # belum bayar, ditulis catatan
            ws.cell(row=r, column=kol, value=rng.choice(["blm","BELUM","nunggak","x"]))
            tunggak_bulan += 1
    ws.cell(row=r, column=17, value=ket)
    bulan_wajib = keluar_di if keluar_di is not None else 12
    fakta.append(dict(nis=nis, kelas=kelas, spp=spp, bulan_wajib=bulan_wajib,
                      bayar=bayar_total, bayar_angka=bayar_angka, tagihan=spp*bulan_wajib))
    r += 1; n_siswa += 1

# baris total di bawah: TU memakai =SUM() Excel, yang MENGABAIKAN teks 'LUNAS'
versi_tu = sum(f["bayar_angka"] for f in fakta)
c = ws.cell(row=r+1, column=2, value="TOTAL DITERIMA (versi TU)"); c.font = Font(bold=True)
c2 = ws.cell(row=r+1, column=16, value=versi_tu); c2.font = Font(bold=True)
wb.save("Rekap SPP Yayasan Bina Cendekia.xlsx")

benar_tagihan = sum(f["tagihan"] for f in fakta)
benar_bayar   = sum(f["bayar"] for f in fakta)
print("siswa            :", n_siswa)
print("total tagihan     :", f"{benar_tagihan:,}".replace(",","."))
print("total diterima    :", f"{benar_bayar:,}".replace(",","."))
print("tunggakan         :", f"{benar_tagihan-benar_bayar:,}".replace(",","."))
print("siswa pindah      :", sum(1 for f in fakta if f["bulan_wajib"] < 12))
print()
print("sel bertuliskan LUNAS :", n_lunas)
print("versi TU (=SUM Excel) :", f"{versi_tu:,}".replace(",","."))
print("selisih yang disembunyikan LUNAS:", f"{benar_bayar-versi_tu:,}".replace(",","."))
